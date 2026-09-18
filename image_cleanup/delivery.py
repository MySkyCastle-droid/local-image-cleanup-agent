"""Portable client package without original files or source metadata."""

from __future__ import annotations

import csv
import json
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageOps

from .workspace import audit_path, load_report


def _thumbnail(source: Path, target: Path, size: tuple[int, int] = (1800, 1200)) -> None:
    with Image.open(source) as opened:
        image = ImageOps.contain(opened.convert("RGB"), size)
        image.save(target, format="JPEG", quality=88)


def _add_presentation(path: Path, rows: list[dict], stage: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    title_slide = deck.slides.add_slide(deck.slide_layouts[6])
    title_box = title_slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(12), Inches(1.2))
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = "Image cleanup delivery"
    paragraph.font.size = Pt(32)
    summary_box = title_slide.shapes.add_textbox(Inches(0.7), Inches(2.4), Inches(12), Inches(1))
    summary_box.text_frame.paragraphs[0].text = f"Completed images: {len(rows)}"
    for row in rows:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.45), Inches(0.2), Inches(12), Inches(0.5))
        paragraph = title.text_frame.paragraphs[0]
        paragraph.text = f"{row['id']}  |  {row['qa_status']}"
        paragraph.font.size = Pt(18)
        comparison = stage / row["comparison"]
        with Image.open(comparison) as opened:
            width, height = opened.size
        max_w, max_h = 12.4, 6.35
        scale = min(max_w / width, max_h / height)
        display_w, display_h = width * scale, height * scale
        slide.shapes.add_picture(
            str(comparison), Inches((13.333 - display_w) / 2),
            Inches(0.78 + (max_h - display_h) / 2),
            width=Inches(display_w), height=Inches(display_h),
        )
    deck.core_properties.author = "Local Image Cleanup Agent"
    deck.core_properties.last_modified_by = ""
    deck.save(path)


def _add_workbook(path: Path, rows: list[dict], totals: dict) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.append(["Category", "Count"])
    for key, value in totals.items():
        overview.append([key, value])
    detail = workbook.create_sheet("Images")
    detail.append(["ID", "Processing", "QA", "QA origin", "Warnings", "Result"])
    for row in rows:
        detail.append([
            row["id"], row["processing_status"], row["qa_status"], row["qa_origin"],
            ";".join(row["warnings"]), row.get("cleaned", ""),
        ])
    workbook.properties.creator = "Local Image Cleanup Agent"
    workbook.properties.lastModifiedBy = ""
    workbook.save(path)


def export_delivery(output: Path) -> Path:
    output = output.resolve()
    payload = load_report(output)
    if not payload["records"]:
        raise ValueError("No processed images available")
    rows: list[dict] = []
    destination = output / "client_delivery.zip"
    with tempfile.TemporaryDirectory(prefix="delivery-", dir=output) as temporary:
        stage = Path(temporary)
        (stage / "results").mkdir()
        (stage / "comparisons").mkdir()
        for record in sorted(payload["records"].values(), key=lambda row: row["id"]):
            accepted = (
                record.get("qa_status") in {"completed_auto", "completed_manual"}
                and record.get("processing_status") in {"cleaned", "no_change"}
            )
            row = {
                "id": record["id"],
                "processing_status": record.get("processing_status", "failed"),
                "qa_status": record.get("qa_status", "needs_review"),
                "qa_origin": record.get("qa_origin", "automatic"),
                "warnings": record.get("warnings", []),
                "cleaned": "",
                "comparison": "",
            }
            if accepted:
                cleaned_source = audit_path(output, record["artifacts"]["cleaned"])
                comparison_source = audit_path(output, record["artifacts"]["comparison"])
                if not cleaned_source.is_file() or not comparison_source.is_file():
                    raise FileNotFoundError(f"Missing audit result for {record['id']}")
                cleaned_name = f"results/{record['id']}{cleaned_source.suffix.lower()}"
                comparison_name = f"comparisons/{record['id']}.jpg"
                with Image.open(cleaned_source) as opened:
                    image = opened.convert("RGBA" if "A" in opened.getbands() else "RGB")
                    suffix = cleaned_source.suffix.lower()
                    if suffix in {".jpg", ".jpeg"}:
                        image.convert("RGB").save(stage / cleaned_name, format="JPEG", quality=95, subsampling=0)
                    elif suffix == ".png":
                        image.save(stage / cleaned_name, format="PNG")
                    else:
                        image.save(stage / cleaned_name, format="WEBP", quality=95)
                _thumbnail(comparison_source, stage / comparison_name)
                row["cleaned"] = cleaned_name
                row["comparison"] = comparison_name
            rows.append(row)

        totals = {
            "total": len(rows),
            "delivered": sum(bool(row["cleaned"]) for row in rows),
            "needs_review": sum(row["qa_status"] == "needs_review" for row in rows),
            "imperfect": sum(row["qa_status"] == "imperfect" for row in rows),
            "failed": sum(row["processing_status"] == "failed" for row in rows),
        }
        (stage / "delivery.json").write_text(
            json.dumps({"summary": totals, "records": rows}, indent=2), encoding="utf-8"
        )
        with (stage / "delivery.csv").open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=("id", "processing_status", "qa_status", "qa_origin", "warnings", "cleaned", "comparison"),
            )
            writer.writeheader()
            for row in rows:
                writer.writerow({**row, "warnings": ";".join(row["warnings"])})
        _add_workbook(stage / "delivery.xlsx", rows, totals)
        _add_presentation(stage / "comparisons.pptx", [row for row in rows if row["cleaned"]], stage)
        temporary_zip = output / ".client_delivery.zip.partial"
        with zipfile.ZipFile(temporary_zip, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(stage.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(stage).as_posix())
        temporary_zip.replace(destination)
    return destination
