from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import numpy as np
from PIL import Image
from openpyxl import load_workbook
from pptx import Presentation

from image_cleanup.delivery import export_delivery
from image_cleanup.engine import process
from image_cleanup.workspace import (
    item_id, load_report, mark_box, process_batch, set_qa, write_report,
)


class WorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        root = Path(self.folder.name)
        self.source = root / "source"
        self.output = root / "audit"
        self.source.mkdir()
        image = Image.new("RGB", (120, 80), (160, 185, 195))
        exif = Image.Exif()
        exif[315] = "private-camera-name"
        image.save(self.source / "alpha.jpg", exif=exif)
        Image.new("RGB", (120, 80), (90, 140, 180)).save(self.source / "beta.png")

    def _run(self, **kwargs):
        return process_batch(
            self.source, self.output, authorized=True, detector="none", method="telea", **kwargs
        )

    def test_mask_changes_only_selected_pixels(self) -> None:
        image = Image.new("RGB", (32, 24), (120, 130, 140))
        result = process(
            image, detector="none", method="telea", manual_boxes=[[4, 5, 10, 8]]
        )
        original = np.asarray(image)
        cleaned = np.asarray(result.cleaned)
        mask = np.asarray(result.mask) > 0
        self.assertTrue(np.array_equal(cleaned[~mask], original[~mask]))
        self.assertEqual(int(mask.sum()), 80)
        self.assertIn("manual_region_requires_review", result.warnings)

    def test_empty_mask_needs_review_and_detected_mask_can_auto_complete(self) -> None:
        empty = self._run(max_new=1)
        self.assertEqual(empty["records"][item_id("alpha.jpg")]["qa_status"], "needs_review")
        self.assertIn("empty_mask", empty["records"][item_id("alpha.jpg")]["warnings"])
        image = Image.new("RGB", (120, 80), (160, 185, 195))
        with patch("image_cleanup.engine.detect_text", side_effect=[
            [([5, 5, 12, 8], 0.99)], []
        ]):
            result = process(image, detector="ocr", method="telea")
        self.assertTrue(result.detected)
        self.assertEqual(result.warnings, [])
        with patch("image_cleanup.engine.detect_text", side_effect=[
            [([5, 5, 12, 8], 0.99)], []
        ]):
            updated = process_batch(
                self.source, self.output, authorized=True, detector="ocr",
                method="telea", only_id=item_id("alpha.jpg"),
            )
        self.assertEqual(updated["records"][item_id("alpha.jpg")]["qa_status"], "completed_auto")

    def test_processing_failure_is_visible_and_not_delivered(self) -> None:
        with patch("image_cleanup.workspace.process", side_effect=RuntimeError("sensitive detail")):
            report = self._run(max_new=1)
        record = report["records"][item_id("alpha.jpg")]
        self.assertEqual(record["processing_status"], "failed")
        self.assertEqual(record["qa_status"], "needs_review")
        self.assertEqual(record["warnings"], ["processing_failed"])
        self.assertNotIn("sensitive detail", (self.output / "report.json").read_text(encoding="utf-8"))
        with zipfile.ZipFile(export_delivery(self.output)) as archive:
            manifest = json.loads(archive.read("delivery.json"))
            self.assertEqual(manifest["summary"]["delivered"], 0)

    def test_resume_review_and_client_delivery(self) -> None:
        first = self._run(max_new=1)
        self.assertEqual(first["summary"]["total"], 1)
        second = self._run()
        self.assertEqual(second["summary"]["total"], 2)
        identifier = item_id("alpha.jpg")
        record = second["records"][identifier]
        self.assertEqual(record["qa_status"], "needs_review")
        self.assertFalse(Path(record["artifacts"]["cleaned"]).is_absolute())

        mark_box(self.output, identifier, [10, 10, 20, 12])
        with self.assertRaises(ValueError):
            set_qa(self.output, identifier, "completed_manual")
        third = self._run(only_id=identifier)
        self.assertEqual(third["records"][identifier]["qa_status"], "needs_review")
        set_qa(self.output, identifier, "completed_manual", "checked locally")
        resumed = self._run()
        self.assertEqual(resumed["records"][identifier]["qa_status"], "completed_manual")
        self.assertEqual(len(resumed["records"][identifier]["qa_history"]), 1)
        set_qa(self.output, item_id("beta.png"), "completed_manual", "no overlay in source")

        package = export_delivery(self.output)
        with zipfile.ZipFile(package) as archive:
            names = set(archive.namelist())
            self.assertIn("comparisons.pptx", names)
            self.assertIn("delivery.xlsx", names)
            self.assertNotIn("alpha.jpg", names)
            self.assertFalse(any("originals/" in name for name in names))
            manifest = json.loads(archive.read("delivery.json"))
            self.assertEqual(manifest["summary"]["delivered"], 2)
            self.assertNotIn("alpha.jpg", archive.read("delivery.json").decode())
            Presentation(io.BytesIO(archive.read("comparisons.pptx")))
            workbook = load_workbook(io.BytesIO(archive.read("delivery.xlsx")), read_only=True)
            self.assertIn("Images", workbook.sheetnames)
            with Image.open(io.BytesIO(archive.read(f"results/{identifier}.jpg"))) as cleaned:
                self.assertFalse(cleaned.getexif())
        with Image.open(self.output / record["artifacts"]["original"]) as original:
            self.assertEqual(original.getexif().get(315), "private-camera-name")

    def test_changed_source_requires_new_review_but_keeps_history(self) -> None:
        self._run()
        identifier = item_id("alpha.jpg")
        set_qa(self.output, identifier, "completed_manual", "accepted")
        Image.new("RGB", (120, 80), (20, 30, 40)).save(self.source / "alpha.jpg")
        self._run()
        record = load_report(self.output)["records"][identifier]
        self.assertEqual(record["qa_status"], "needs_review")
        self.assertEqual(record["manual_qa"]["status"], "completed_manual")
        self.assertEqual(len(record["qa_history"]), 1)

    def test_source_and_output_must_be_separate(self) -> None:
        with self.assertRaises(ValueError):
            process_batch(
                self.source, self.source / "out", authorized=True,
                detector="none", method="telea",
            )

    def test_removed_source_is_not_delivered(self) -> None:
        self._run()
        set_qa(self.output, item_id("beta.png"), "completed_manual")
        (self.source / "alpha.jpg").unlink()
        report = self._run()
        self.assertEqual(report["records"][item_id("alpha.jpg")]["processing_status"], "missing_source")
        self.assertEqual(report["summary"]["failed"], 1)
        with zipfile.ZipFile(export_delivery(self.output)) as archive:
            manifest = json.loads(archive.read("delivery.json"))
            self.assertEqual(manifest["summary"]["delivered"], 1)
        with self.assertRaises(ValueError):
            set_qa(self.output, item_id("alpha.jpg"), "completed_manual")
        (self.source / "beta.png").unlink()
        report = self._run()
        self.assertEqual(report["summary"]["failed"], 2)
        with zipfile.ZipFile(export_delivery(self.output)) as archive:
            manifest = json.loads(archive.read("delivery.json"))
            self.assertEqual(manifest["summary"]["delivered"], 0)

    def test_tampered_audit_path_cannot_escape_output(self) -> None:
        self._run()
        set_qa(self.output, item_id("alpha.jpg"), "completed_manual")
        payload = load_report(self.output)
        payload["records"][item_id("alpha.jpg")]["artifacts"]["cleaned"] = "../source/alpha.jpg"
        write_report(self.output, payload)
        with self.assertRaises(ValueError):
            export_delivery(self.output)

    def test_transparent_png_keeps_alpha_without_metadata(self) -> None:
        transparent = Image.new("RGBA", (32, 24), (120, 130, 140, 90))
        transparent.save(self.source / "clear.png")
        self._run()
        identifier = item_id("clear.png")
        set_qa(self.output, identifier, "completed_manual")
        report = load_report(self.output)
        with Image.open(self.output / report["records"][identifier]["artifacts"]["cleaned"]) as image:
            self.assertEqual(image.getpixel((0, 0))[3], 90)
        with zipfile.ZipFile(export_delivery(self.output)) as archive:
            with Image.open(io.BytesIO(archive.read(f"results/{identifier}.png"))) as image:
                self.assertEqual(image.getpixel((0, 0))[3], 90)
                self.assertFalse(image.getexif())


if __name__ == "__main__":
    unittest.main()
